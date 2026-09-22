#!/usr/bin/env python3
"""
英語学習アプリ 問題データ変換スクリプト

問題データの Excel を読み、アプリが読み込む JSON（データパック）を作る。
あわせて、シート間の整合性（まとめ 15-11 など）を点検する。

使い方（アプリのフォルダで実行する）
  py tools/build_data.py 問題データ.xlsx

出力
  data/index.json       パックの一覧と版数
  data/common.json      文法ポイント・単語・不規則動詞・熟語・数詞・聞き取りゆれ
  data/pack_0001.json   S0001〜S0050 の文に関するデータ（50文ごと）

必要なもの
  Python 3.9 以上と openpyxl（py -m pip install openpyxl）
"""

import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path

PACK_SIZE = 50
LABELS = ['主語', '動詞', '助動詞', '目的語', '間接目的語', '補語']

# シートの先頭番号 → （データ名, {Excelの列名: JSONのキー}）
SHEETS = {
    '01_': ('sentences', {'文ID': 'id', '英文': 'en', '日本語訳': 'ja', '教材名': 'textbook',
                          'Lesson/Unit': 'unit', '学年': 'grade', '文法ポイントID': 'grammarId',
                          '難易度(1-5)': 'difficulty', '備考': 'note', '読み上げ用日本語': 'jaTts'}),
    '02_': ('grammar', {'文法ポイントID': 'id', '文法ポイント名': 'name', '学年区分': 'grade',
                        '説明': 'desc', '備考': 'note'}),
    '03_': ('words', {'単語ID': 'id', '単語(見出し語)': 'word', '品詞': 'pos', '正解の意味': 'meaning',
                      'ダミー意味1': 'dummy1', 'ダミー意味2': 'dummy2', 'ダミー意味3': 'dummy3',
                      '発音記号': 'ipa', '備考': 'note'}),
    '04_': ('sentenceWords', {'文ID': 'sid', '単語ID': 'wid', '出題順': 'order', '備考': 'note'}),
    '05_': ('structure', {'文ID': 'sid', '語順番号': 'no', '語': 'word', '文要素ラベル': 'label',
                          '出題順': 'order', '備考': 'note'}),
    '06_': ('structureChoice', {'文ID': 'sid', '正解の和訳': 'correct',
                                '誤答の和訳1': 'wrong1', '誤読の意図1': 'intent1',
                                '誤答の和訳2': 'wrong2', '誤読の意図2': 'intent2',
                                '誤答の和訳3': 'wrong3', '誤読の意図3': 'intent3'}),
    '07_': ('cloze', {'文ID': 'sid', '空欄位置(語順番号)': 'blankNo', '正解': 'answer',
                      'ダミー選択肢1': 'dummy1', 'ダミー選択肢2': 'dummy2', 'ダミー選択肢3': 'dummy3',
                      '出題意図/文法ポイント': 'intent', '誤答類型': 'errorType', '備考': 'note'}),
    '08_': ('order', {'文ID': 'sid', '語順番号': 'no', '語': 'word', '品詞': 'pos',
                      '不要語フラグ(1=不要語)': 'extra', '備考': 'note'}),
    '09_': ('patterns', {'文ID': 'sid', 'パターン番号': 'no', '許容する英文表記': 'text', '備考': 'note'}),
    '10_': ('irregularVerbs', {'動詞ID': 'id', '原形': 'base', '過去形': 'past', '過去分詞': 'pp',
                               '意味': 'meaning', '変化型': 'type', '備考': 'note'}),
    '11_': ('idioms', {'熟語ID': 'id', '熟語': 'idiom', '意味': 'meaning', '例文(英)': 'exEn',
                       '例文(日)': 'exJa', 'ダミー意味1': 'dummy1', 'ダミー意味2': 'dummy2',
                       'ダミー意味3': 'dummy3', '関連文ID': 'sid', '備考': 'note'}),
    '12_': ('numbers', {'数詞ID': 'id', '区分': 'category', '表記': 'written', '英語表記': 'english',
                        '読み方の注意': 'readingNote', '備考': 'note'}),
    '13_': ('mishear', {'正しい語': 'correct', '聞き取られた語': 'heard', '対立の種類': 'type',
                        '確認状況': 'status', '備考': 'note'}),
}
PER_SENTENCE = ['sentences', 'sentenceWords', 'structure', 'structureChoice', 'cloze', 'order', 'patterns']
COMMON = ['grammar', 'words', 'irregularVerbs', 'idioms', 'numbers', 'mishear']


def clean(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def read_workbook(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit('openpyxl が必要です。py -m pip install openpyxl を実行してください。')
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data = {name: [] for name, _ in SHEETS.values()}
    for ws in wb.worksheets:
        spec = next((v for k, v in SHEETS.items() if ws.title.startswith(k)), None)
        if not spec:
            continue
        name, colmap = spec
        heads = [clean(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        unknown = [h for h in heads if h and h not in colmap]
        if unknown:
            print(f'  注意：{ws.title} の列 {unknown} は取り込みません')
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = {}
            for h, v in zip(heads, row):
                if h in colmap:
                    v = clean(v)
                    if v is not None:
                        rec[colmap[h]] = v
            if rec:
                # 文法ポイントは「G003;G007」のように複数書ける
                if name == 'sentences' and 'grammarId' in rec:
                    rec['grammarIds'] = [g for g in re.split(r'[;；,、\s]+', str(rec.pop('grammarId'))) if g]
                data[name].append(rec)
    return data


def words_of(sentence):
    """英文を語に分ける（点検用）。記号は除き、アポストロフィは残す。"""
    t = sentence.replace('\u2019', "'")
    return [w for w in re.split(r"[^A-Za-z0-9']+", t) if w]


def validate(d):
    errors, warns = [], []
    E, W = errors.append, warns.append
    sids = {}
    for s in d['sentences']:
        sid = s.get('id', '')
        if not re.fullmatch(r'S\d{4}', str(sid)):
            E(f'01：文ID「{sid}」は S0001 形式ではありません')
            continue
        if sid in sids:
            E(f'01：文ID {sid} が重複しています')
        sids[sid] = s
        for k, label in (('en', '英文'), ('ja', '日本語訳')):
            if not s.get(k):
                E(f'01：{sid} の{label}が空です')
    gids = {g.get('id') for g in d['grammar']}
    wids = {w.get('id'): w for w in d['words']}
    for s in sids.values():
        for g in s.get('grammarIds', []):
            if g not in gids:
                E(f'01：{s["id"]} の文法ポイントID {g} が 02 にありません')

    for w in d['words']:
        for k, label in (('word', '単語'), ('pos', '品詞'), ('meaning', '正解の意味'),
                         ('dummy1', 'ダミー意味1'), ('dummy2', 'ダミー意味2'), ('dummy3', 'ダミー意味3')):
            if not w.get(k):
                E(f'03：{w.get("id")} の{label}が空です')
        opts = [w.get('meaning'), w.get('dummy1'), w.get('dummy2'), w.get('dummy3')]
        if len(set(opts)) != len(opts):
            E(f'03：{w.get("id")} の選択肢に同じものがあります')

    def check_sid(sheet, r):
        if r.get('sid') not in sids:
            E(f'{sheet}：文ID {r.get("sid")} が 01 にありません')
            return False
        return True

    for r in d['sentenceWords']:
        if check_sid('04', r) and r.get('wid') not in wids:
            E(f'04：{r["sid"]} の単語ID {r.get("wid")} が 03 にありません')

    by = lambda rows: {sid: sorted([r for r in rows if r.get('sid') == sid], key=lambda r: r.get('no', 0))
                       for sid in sids}
    order = by([r for r in d['order'] if check_sid('08', r)])
    struct = by([r for r in d['structure'] if check_sid('05', r)])

    for sid, rows in order.items():
        if not rows:
            W(f'08：{sid} の語順並べ替えのデータがありません')
            continue
        main = [r for r in rows if r.get('extra') != 1]
        extra = [r for r in rows if r.get('extra') == 1]
        if len(extra) != 1:
            W(f'08：{sid} の不要語が {len(extra)} 語です（1語の想定）')
        got = [str(r.get('word', '')).lower() for r in main]
        want = [w.lower() for w in words_of(sids[sid]['en'])]
        if got != want:
            E(f'08：{sid} の語を並べても英文になりません\n      08：{" ".join(got)}\n      01：{" ".join(want)}')
        nos = [r.get('no') for r in main]
        if nos != list(range(1, len(main) + 1)):
            E(f'08：{sid} の語順番号が 1 からの連番になっていません：{nos}')

    for sid, rows in struct.items():
        if not rows:
            W(f'05：{sid} の構造把握のデータがありません')
            continue
        main = {r.get('no'): str(r.get('word', '')) for r in order.get(sid, []) if r.get('extra') != 1}
        for r in rows:
            if r.get('label') and r['label'] not in LABELS:
                E(f'05：{sid} の「{r.get("word")}」のラベル「{r["label"]}」は6種類のいずれでもありません')
            if main and main.get(r.get('no')) != str(r.get('word', '')):
                E(f'05：{sid} 語順番号 {r.get("no")} の語「{r.get("word")}」が 08 の「{main.get(r.get("no"))}」と一致しません')
        if not any(r.get('label') for r in rows):
            W(f'05：{sid} に文要素ラベルが1つもありません')

    for r in d['cloze']:
        if not check_sid('07', r):
            continue
        main = {x.get('no'): str(x.get('word', '')) for x in order.get(r['sid'], []) if x.get('extra') != 1}
        if main and main.get(r.get('blankNo')) != str(r.get('answer', '')):
            E(f'07：{r["sid"]} の空欄位置 {r.get("blankNo")} の語は「{main.get(r.get("blankNo"))}」で、正解「{r.get("answer")}」と一致しません')
        opts = [r.get('answer'), r.get('dummy1'), r.get('dummy2'), r.get('dummy3')]
        if None in opts:
            E(f'07：{r["sid"]} の選択肢が4つそろっていません')

    for r in d['structureChoice']:
        if check_sid('06', r):
            if not all(r.get(k) for k in ('correct', 'wrong1', 'wrong2', 'wrong3')):
                E(f'06：{r["sid"]} の和訳の選択肢が4つそろっていません')
            elif r['correct'] != sids[r['sid']]['ja']:
                W(f'06：{r["sid"]} の正解の和訳が 01 の日本語訳と一致しません')

    for r in d['patterns']:
        check_sid('09', r)
    for r in d['idioms']:
        if r.get('sid') and r['sid'] not in sids:
            W(f'11：{r.get("id")} の関連文ID {r["sid"]} が 01 にありません')
    for sid in sids:
        for sheet, rows in (('04', d['sentenceWords']), ('06', d['structureChoice']), ('07', d['cloze'])):
            if not any(r.get('sid') == sid for r in rows):
                W(f'{sheet}：{sid} のデータがありません')
    return errors, warns


def digest(obj):
    return hashlib.sha256(json.dumps(obj, ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()[:12]


def main():
    if len(sys.argv) < 2:
        sys.exit('使い方：py tools/build_data.py 問題データ.xlsx')
    src = sys.argv[1]
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path('data')
    print(f'読み込み：{src}')
    d = read_workbook(src)
    print('  ' + '、'.join(f'{k} {len(v)}' for k, v in d.items()))

    errors, warns = validate(d)
    for w in warns:
        print(f'  注意：{w}')
    if errors:
        print()
        for e in errors:
            print(f'  誤り：{e}')
        sys.exit(f'\n{len(errors)} 件の誤りがあります。Excel を直してから、もう一度実行してください。')

    out.mkdir(parents=True, exist_ok=True)
    packs = {}
    for s in d['sentences']:
        n = (int(s['id'][1:]) - 1) // PACK_SIZE
        packs.setdefault(n, set()).add(s['id'])
    index = {'generatedAt': datetime.now().isoformat(timespec='seconds'), 'packs': []}
    written = set()
    for n in sorted(packs):
        ids = packs[n]
        body = {k: [r for r in d[k] if (r.get('id') if k == 'sentences' else r.get('sid')) in ids]
                for k in PER_SENTENCE}
        name = f'pack_{n + 1:04d}.json'
        lo, hi = n * PACK_SIZE + 1, (n + 1) * PACK_SIZE
        (out / name).write_text(json.dumps(body, ensure_ascii=False, indent=1), encoding='utf-8')
        written.add(name)
        index['packs'].append({'file': name, 'range': [f'S{lo:04d}', f'S{hi:04d}'],
                               'sentences': len(ids), 'version': digest(body)})
    common = {k: d[k] for k in COMMON}
    (out / 'common.json').write_text(json.dumps(common, ensure_ascii=False, indent=1), encoding='utf-8')
    index['common'] = {'file': 'common.json', 'version': digest(common)}
    (out / 'index.json').write_text(json.dumps(index, ensure_ascii=False, indent=1), encoding='utf-8')

    stale = sorted(p.name for p in out.glob('pack_*.json') if p.name not in written)
    for p in index['packs']:
        print(f'  {p["file"]}  {p["range"][0]}〜{p["range"][1]}  {p["sentences"]}文  版 {p["version"]}')
    print(f'  common.json  版 {index["common"]["version"]}')
    if stale:
        print(f'  注意：使われていないパックが残っています：{", ".join(stale)}')
    print(f'\n完了：{out.resolve()}')


if __name__ == '__main__':
    main()
