#!/usr/bin/env python3
"""
英語学習アプリ 音声生成スクリプト（Google Cloud Text-to-Speech）

問題データの Excel を読み、文と単語の音声を MP3 で生成する。
すでに生成済みで内容が変わっていない音声は作り直さない。

使い方（Windows のコマンドプロンプトの例）
  py generate_audio.py 問題データ.xlsx                足りない音声だけ生成する
  py generate_audio.py 問題データ.xlsx --dry-run      生成する予定の一覧だけ表示する
  py generate_audio.py 問題データ.xlsx --only S0001   指定したIDだけ生成する（複数可）
  py generate_audio.py 問題データ.xlsx --force        すべて作り直す
  py generate_audio.py --list-voices en-US            使える声の一覧を表示する
  py generate_audio.py 問題データ.xlsx --compare      ビットレートの聞き比べ用サンプルを作る
  py generate_audio.py 問題データ.xlsx --compare S0004 S0010   文を指定して作る

必要なもの
  Python 3.9 以上と openpyxl・lameenc（py -m pip install openpyxl lameenc）
  Text-to-Speech API を有効にした Google Cloud の API キー
    環境変数 GOOGLE_TTS_API_KEY に入れるか、このスクリプトと同じフォルダの
    tts_api_key.txt に1行で書いておく。
    tts_api_key.txt は GitHub などに上げないこと。

出力
  audio/male/S0001_ja.mp3 など（声の種類ごとにフォルダを分ける）
  audio/silence_500ms.mp3       連結時に文の間に挟む無音
  audio/manifest.json            各ファイルの正確な長さ（フレーム数から計算）と、
                                 作り直しの判定に使う内容のハッシュ

音声の形式
  Google からは無圧縮（LINEAR16）で受け取り、手元で固定ビットレートの MP3 に変換する。
  Google が直接返す MP3 は 32kbps 固定のため。圧縮は1回だけなので劣化が重ならない。
  ビットレートは MP3_BITRATE_KBPS で決める。--compare で聞き比べてから決めること。

MP3 の後処理
  変換した MP3 から ID3 タグと先頭の管理用フレーム（Xing / Info / VBRI）を取り除き、
  音声のフレームだけにする。これを残したまま連結すると、ブラウザが最初のファイルの
  長さを全体の長さと取り違え、再生位置と表示中の文がずれるため。
"""

import argparse
import base64
import hashlib
import io
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
import wave
from datetime import datetime
from pathlib import Path

# ======================================================================
# 設定（必要に応じてここを書き換える）
# ======================================================================

OUTPUT_DIR = 'audio'
SAMPLE_RATE = 24000          # すべての音声を同じサンプリング周波数で作る（連結の前提）。Neural2 の声は 24kHz
MP3_BITRATE_KBPS = 128       # 本番の生成に使うビットレート。--compare で聞き比べて決める
COMPARE_BITRATES = [32, 64, 128]
COMPARE_DEFAULT_IDS = ['S0004']   # She / violin / five（sh・v・f の音）を含む文

# 声の種類 → 言語ごとの声の名前。--list-voices で使える名前を確認できる。
VOICES = {
    'male':   {'en': 'en-US-Neural2-D', 'ja': 'ja-JP-Neural2-C'},
    'female': {'en': 'en-US-Neural2-F', 'ja': 'ja-JP-Neural2-B'},
}

# 読み上げ速度（まとめ 10-1 の既定値）
EN_SPEEDS = {'normal': 0.85, 'slow': 0.65, 'fast': 1.25}
JA_SPEED = 1.25
WORD_SPEED = 0.85            # 単語の音声（意味理解ステップで再生）

SILENCES_MS = [500]          # 連結用の無音。長い間隔はアプリ側で繰り返して作る

SENTENCE_SHEET_PREFIX = '01_'
WORD_SHEET_PREFIX = '03_'
JA_TTS_COLUMN = '読み上げ用日本語'   # 01_文マスタにこの列があれば、日本語訳より優先して読む

API_BASE = 'https://texttospeech.googleapis.com/v1'
LANG_CODE = {'en': 'en-US', 'ja': 'ja-JP'}

# ======================================================================
# API キー
# ======================================================================

def load_api_key():
    key = os.environ.get('GOOGLE_TTS_API_KEY', '').strip()
    if key:
        return key
    f = Path(__file__).with_name('tts_api_key.txt')
    if f.exists():
        key = f.read_text(encoding='utf-8').strip()
        if key:
            return key
    sys.exit('API キーが見つかりません。環境変数 GOOGLE_TTS_API_KEY を設定するか、\n'
             f'{f} に API キーを1行で書いてください。')

# ======================================================================
# Google Cloud TTS の呼び出し
# ======================================================================

def api_request(url, body=None, retries=4):
    data = json.dumps(body).encode('utf-8') if body is not None else None
    req = urllib.request.Request(url, data=data, method='POST' if body is not None else 'GET',
                                 headers={'Content-Type': 'application/json; charset=utf-8'})
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(req, timeout=60) as res:
                return json.loads(res.read().decode('utf-8'))
        except urllib.error.HTTPError as e:
            detail = e.read().decode('utf-8', 'replace')
            if e.code in (429, 500, 502, 503, 504) and attempt < retries - 1:
                wait = 2 ** attempt
                print(f'    一時的なエラー（{e.code}）。{wait}秒後に再試行します')
                time.sleep(wait)
                continue
            try:
                msg = json.loads(detail)['error']['message']
            except Exception:
                msg = detail[:500]
            raise RuntimeError(f'API エラー {e.code}: {msg}') from None
        except urllib.error.URLError as e:
            if attempt < retries - 1:
                time.sleep(2 ** attempt)
                continue
            raise RuntimeError(f'通信エラー: {e.reason}') from None


def list_voices(key, lang):
    res = api_request(f'{API_BASE}/voices?languageCode={lang}&key={key}')
    return res.get('voices', [])


def synthesize_pcm(key, voice_name, lang, speed, text):
    """Google から無圧縮（LINEAR16・WAV）で受け取り、16bit モノラルの PCM を返す。"""
    body = {
        'input': {'text': text},
        'voice': {'languageCode': LANG_CODE[lang], 'name': voice_name},
        'audioConfig': {'audioEncoding': 'LINEAR16', 'speakingRate': speed, 'sampleRateHertz': SAMPLE_RATE},
    }
    res = api_request(f'{API_BASE}/text:synthesize?key={key}', body)
    return wav_to_pcm(base64.b64decode(res['audioContent']))


def wav_to_pcm(wav_bytes):
    with wave.open(io.BytesIO(wav_bytes), 'rb') as w:
        ch, width, sr = w.getnchannels(), w.getsampwidth(), w.getframerate()
        pcm = w.readframes(w.getnframes())
    if (ch, width, sr) != (1, 2, SAMPLE_RATE):
        raise ValueError(f'想定外の形式です（{ch}ch・{width * 8}bit・{sr}Hz）。'
                         f'モノラル・16bit・{SAMPLE_RATE}Hz を想定しています。')
    return pcm


def pcm_to_wav(pcm):
    buf = io.BytesIO()
    with wave.open(buf, 'wb') as w:
        w.setnchannels(1); w.setsampwidth(2); w.setframerate(SAMPLE_RATE)
        w.writeframes(pcm)
    return buf.getvalue()


def encode_mp3(pcm, kbps):
    """固定ビットレートの MP3 に変換する。出力の周波数も固定する
    （指定しないと、低いビットレートのとき LAME が周波数を勝手に下げるため）。"""
    try:
        import lameenc
    except ImportError:
        sys.exit('lameenc が必要です。py -m pip install lameenc を実行してください。')
    e = lameenc.Encoder()
    e.set_bit_rate(kbps)
    e.set_in_sample_rate(SAMPLE_RATE)
    e.set_out_sample_rate(SAMPLE_RATE)
    e.set_channels(1)
    e.set_quality(2)          # 2 = 高品質
    return e.encode(pcm) + e.flush()


def silence_pcm(ms):
    return b'\x00\x00' * (SAMPLE_RATE * ms // 1000)

# ======================================================================
# MP3 の解析と後処理
# ======================================================================

_BITRATE = {
    1: [0, 32, 40, 48, 56, 64, 80, 96, 112, 128, 160, 192, 224, 256, 320],   # MPEG-1 Layer III
    2: [0, 8, 16, 24, 32, 40, 48, 56, 64, 80, 96, 112, 128, 144, 160],       # MPEG-2/2.5 Layer III
}
_SAMPLE_RATE = {1: [44100, 48000, 32000], 2: [22050, 24000, 16000], 25: [11025, 12000, 8000]}


def _frame_header(b, i):
    if i + 4 > len(b):
        return None
    h = int.from_bytes(b[i:i + 4], 'big')
    if (h >> 21) & 0x7FF != 0x7FF:
        return None
    ver_bits, layer_bits = (h >> 19) & 3, (h >> 17) & 3
    br_idx, sr_idx, pad, ch = (h >> 12) & 0xF, (h >> 10) & 3, (h >> 9) & 1, (h >> 6) & 3
    if ver_bits == 1 or layer_bits != 1 or br_idx in (0, 15) or sr_idx == 3:
        return None                                       # Layer III 以外・不正値
    ver = {3: 1, 2: 2, 0: 25}[ver_bits]
    br = _BITRATE[1 if ver == 1 else 2][br_idx] * 1000
    sr = _SAMPLE_RATE[ver][sr_idx]
    return {
        'ver': ver, 'bitrate': br, 'sr': sr, 'ch': ch,
        'spf': 1152 if ver == 1 else 576,
        'len': (144 if ver == 1 else 72) * br // sr + pad,
    }


def clean_mp3(raw):
    """ID3 タグと先頭の管理用フレームを除き、音声フレームだけを返す。"""
    b = bytes(raw)
    start, end = 0, len(b)
    if b[:3] == b'ID3' and len(b) >= 10:                  # ID3v2
        size = (b[6] << 21) | (b[7] << 14) | (b[8] << 7) | b[9]
        start = 10 + size + (10 if b[5] & 0x10 else 0)
    if end - start >= 128 and b[end - 128:end - 125] == b'TAG':   # ID3v1
        end -= 128

    frames, i = [], start
    while i < end:
        hd = _frame_header(b, i)
        # 次のフレームの頭も正しい場合だけフレームとして認める（誤同期の防止）
        if hd and i + hd['len'] <= end and (i + hd['len'] == end or _frame_header(b, i + hd['len'])):
            frames.append((i, hd))
            i += hd['len']
        else:
            i += 1
    if not frames:
        raise ValueError('MP3 のフレームが見つかりません')

    first_off, first_hd = frames[0]
    first = b[first_off:first_off + first_hd['len']]
    header_frame = any(tag in first for tag in (b'Xing', b'Info', b'VBRI'))
    if header_frame:
        frames = frames[1:]

    out = b''.join(b[o:o + h['len']] for o, h in frames)
    rates = {h['bitrate'] for _, h in frames}
    srs = {h['sr'] for _, h in frames}
    chs = {h['ch'] for _, h in frames}
    if len(srs) != 1:
        raise ValueError(f'サンプリング周波数が混在しています: {srs}')
    sr, spf = frames[0][1]['sr'], frames[0][1]['spf']
    samples = len(frames) * spf
    return out, {
        'frames': len(frames),
        'samples': samples,
        'sampleRate': sr,
        'mpegVersion': frames[0][1]['ver'],
        'channelMode': frames[0][1]['ch'],
        'cbr': len(rates) == 1,
        'bitrateKbps': sorted(r // 1000 for r in rates),
        'durationMs': round(samples * 1000 / sr, 2),
        'bytes': len(out),
        'removedHeaderFrame': header_frame,
        'mixedChannelMode': len(chs) != 1,
    }

# ======================================================================
# Excel の読み込み
# ======================================================================

def _sheet(wb, prefix):
    for ws in wb.worksheets:
        if ws.title.startswith(prefix):
            return ws
    sys.exit(f'「{prefix}」で始まるシートが見つかりません。')


def _rows(ws):
    heads = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and any(v is not None and str(v).strip() for v in row):
            yield {h: v for h, v in zip(heads, row) if h}


def ja_for_tts(text):
    """日本語訳の（　）内は読み上げない。"""
    t = re.sub(r'（[^（）]*）|\([^()]*\)', '', text or '')
    return re.sub(r'\s+', ' ', t).strip()


def load_items(xlsx_path):
    try:
        import openpyxl
    except ImportError:
        sys.exit('openpyxl が必要です。py -m pip install openpyxl を実行してください。')
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    sentences = []
    for r in _rows(_sheet(wb, SENTENCE_SHEET_PREFIX)):
        sid = str(r.get('文ID') or '').strip()
        if not sid:
            continue
        if not re.fullmatch(r'S\d{4}', sid):
            print(f'  注意：文ID「{sid}」は S0001 形式ではありません。読み飛ばします。')
            continue
        en = str(r.get('英文') or '').strip()
        ja_src = str(r.get(JA_TTS_COLUMN) or '').strip() or str(r.get('日本語訳') or '').strip()
        if not en or not ja_src:
            print(f'  注意：{sid} は英文か日本語訳が空です。読み飛ばします。')
            continue
        sentences.append({'id': sid, 'en': en, 'ja': ja_for_tts(ja_src)})

    words = []
    for r in _rows(_sheet(wb, WORD_SHEET_PREFIX)):
        wid = str(r.get('単語ID') or '').strip()
        w = str(r.get('単語(見出し語)') or '').strip()
        if not wid or not w:
            continue
        if not re.fullmatch(r'W\d{4}', wid):
            print(f'  注意：単語ID「{wid}」は W0001 形式ではありません。読み飛ばします。')
            continue
        words.append({'id': wid, 'word': w})

    for label, items in (('文ID', sentences), ('単語ID', words)):
        seen = set()
        for it in items:
            if it['id'] in seen:
                sys.exit(f'{label}「{it["id"]}」が重複しています。')
            seen.add(it['id'])
    return sentences, words

# ======================================================================
# 生成する音声の一覧
# ======================================================================

def build_jobs(sentences, words):
    jobs = []
    for voice_kind, names in VOICES.items():
        for s in sentences:
            jobs.append({'path': f'{voice_kind}/{s["id"]}_ja.mp3', 'id': s['id'], 'lang': 'ja',
                         'voice': names['ja'], 'speed': JA_SPEED, 'text': s['ja']})
            for speed_name, rate in EN_SPEEDS.items():
                jobs.append({'path': f'{voice_kind}/{s["id"]}_en_{speed_name}.mp3', 'id': s['id'], 'lang': 'en',
                             'voice': names['en'], 'speed': rate, 'text': s['en']})
        for w in words:
            jobs.append({'path': f'{voice_kind}/{w["id"]}_en.mp3', 'id': w['id'], 'lang': 'en',
                         'voice': names['en'], 'speed': WORD_SPEED, 'text': w['word']})
    for ms in SILENCES_MS:
        jobs.append({'path': f'silence_{ms}ms.mp3', 'id': f'silence_{ms}ms', 'lang': None,
                     'voice': '-', 'speed': 1.0, 'silence_ms': ms})
    for j in jobs:
        src = json.dumps([j.get('text'), j.get('silence_ms'), j['voice'], j['speed'], SAMPLE_RATE,
                          'LINEAR16', MP3_BITRATE_KBPS], ensure_ascii=False)
        j['hash'] = hashlib.sha256(src.encode('utf-8')).hexdigest()[:16]
    return jobs

# ======================================================================
# メイン
# ======================================================================

def cmd_list_voices(lang):
    key = load_api_key()
    voices = list_voices(key, lang)
    if not voices:
        print(f'{lang} の声が見つかりませんでした。')
        return
    for v in sorted(voices, key=lambda v: v['name']):
        g = {'MALE': '男性', 'FEMALE': '女性'}.get(v.get('ssmlGender'), v.get('ssmlGender', ''))
        print(f'{v["name"]:32} {g}')


COMPARE_HTML = """<!DOCTYPE html>
<html lang="ja"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>音質の聞き比べ</title>
<style>
 :root{{color-scheme:light dark;--bg:#f4f2ec;--card:#fffdf8;--line:#ddd8cc;--text:#22201c;--muted:#6f6a5f}}
 @media (prefers-color-scheme:dark){{:root{{--bg:#1b1a17;--card:#24221e;--line:#3b3831;--text:#e9e5db;--muted:#989285}}}}
 body{{margin:0;background:var(--bg);color:var(--text);font-family:"Hiragino Sans","Noto Sans JP",system-ui,sans-serif;line-height:1.6}}
 main{{max-width:620px;margin:0 auto;padding:24px 18px 60px}}
 h1{{font-size:1.2rem;margin:0 0 6px}} .note{{color:var(--muted);font-size:.875rem}}
 section{{background:var(--card);border:.5px solid var(--line);border-radius:14px;padding:16px;margin:16px 0}}
 h2{{font-size:.95rem;margin:0 0 4px}} .en{{font-family:Georgia,serif;font-size:1.15rem;margin:0 0 10px}}
 .row{{display:flex;align-items:center;gap:12px;padding:8px 0;border-top:.5px solid var(--line)}}
 .row:first-of-type{{border-top:0}} .lbl{{width:1.6em;font-weight:600}} audio{{flex:1;min-width:0}}
 .ans{{color:var(--muted);font-size:.8125rem;min-width:6em;text-align:right;visibility:hidden}}
 .show .ans{{visibility:visible}}
 button{{font:inherit;color:var(--text);background:transparent;border:.5px solid var(--line);border-radius:10px;padding:10px 16px;cursor:pointer}}
</style></head><body><main>
<h1>音質の聞き比べ</h1>
<p class="note">イヤホンで聞いてください。A〜D は、元の無圧縮の音と、32・64・128kbps の MP3 を並べ替えたものです。
She・violin・five の s・sh・v・f の音に注目して、違いが分かるか、どれが元の音かを当ててから「答えを見る」を押してください。</p>
{sections}
<button onclick="document.body.classList.toggle('show')">答えを見る／隠す</button>
</main></body></html>"""


def cmd_compare(args, sentences):
    ids = args.compare or COMPARE_DEFAULT_IDS
    by_id = {s['id']: s for s in sentences}
    missing = [i for i in ids if i not in by_id]
    if missing:
        sys.exit(f'見つからない文ID：{", ".join(missing)}')
    key = load_api_key()
    out = Path(args.out) / 'compare'
    out.mkdir(parents=True, exist_ok=True)
    rnd = __import__('random').Random()
    sections = []
    for sid in ids:
        text = by_id[sid]['en']
        for voice_kind, names in VOICES.items():
            print(f'{sid} {voice_kind}（{names["en"]}）：{text}')
            pcm = synthesize_pcm(key, names['en'], 'en', EN_SPEEDS['normal'], text)
            base = f'{voice_kind}_{sid}'
            files = [(f'{base}_original.wav', '元の無圧縮')]
            (out / files[0][0]).write_bytes(pcm_to_wav(pcm))
            for kbps in COMPARE_BITRATES:
                data, info = clean_mp3(encode_mp3(pcm, kbps))
                name = f'{base}_{kbps}kbps.mp3'
                (out / name).write_bytes(data)
                files.append((name, f'{kbps}kbps'))
                print(f'  {name}  {info["sampleRate"]}Hz・{info["bitrateKbps"][0]}kbps・{info["bytes"] // 1024}KB')
            rnd.shuffle(files)
            rows = ''.join(
                f'<div class="row"><span class="lbl">{"ABCD"[i]}</span>'
                f'<audio controls preload="auto" src="{name}"></audio><span class="ans">{label}</span></div>'
                for i, (name, label) in enumerate(files))
            sections.append(f'<section><h2>{sid}・{"男性" if voice_kind == "male" else "女性"}の声</h2>'
                            f'<p class="en">{text}</p>{rows}</section>')
    html_path = out / 'index.html'
    html_path.write_text(COMPARE_HTML.format(sections=''.join(sections)), encoding='utf-8')
    print(f'\n聞き比べ用のページ：{html_path.resolve()}')
    print('ブラウザで開いて、イヤホンで聞いてください。')


def main():
    ap = argparse.ArgumentParser(description='問題データの Excel から音声を生成します。')
    ap.add_argument('xlsx', nargs='?', help='問題データの Excel ファイル')
    ap.add_argument('--out', default=OUTPUT_DIR, help=f'出力先フォルダ（既定：{OUTPUT_DIR}）')
    ap.add_argument('--only', nargs='+', metavar='ID', help='指定した文ID・単語IDだけ生成する')
    ap.add_argument('--force', action='store_true', help='生成済みの音声も作り直す')
    ap.add_argument('--dry-run', action='store_true', help='生成する予定の一覧だけ表示する')
    ap.add_argument('--list-voices', metavar='LANG', help='使える声の一覧を表示する（例：en-US、ja-JP）')
    ap.add_argument('--compare', nargs='*', metavar='ID',
                    help=f'ビットレート（{"・".join(map(str, COMPARE_BITRATES))}kbps）の聞き比べ用サンプルを作る'
                         f'（文IDを省略すると {" ".join(COMPARE_DEFAULT_IDS)}）')
    args = ap.parse_args()

    if args.list_voices:
        cmd_list_voices(args.list_voices)
        return
    if not args.xlsx:
        ap.error('Excel ファイルを指定してください。')

    print(f'読み込み：{args.xlsx}')
    sentences, words = load_items(args.xlsx)
    print(f'  文 {len(sentences)} 件、単語 {len(words)} 件')
    if args.compare is not None:
        cmd_compare(args, sentences)
        return

    out_dir = Path(args.out)
    manifest_path = out_dir / 'manifest.json'
    manifest = {'files': {}}
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        manifest.setdefault('files', {})

    jobs = build_jobs(sentences, words)
    if args.only:
        wanted = set(args.only)
        jobs = [j for j in jobs if j['id'] in wanted]
        missing = wanted - {j['id'] for j in jobs}
        if missing:
            print(f'  注意：見つからないID：{", ".join(sorted(missing))}')

    todo = []
    for j in jobs:
        rec = manifest['files'].get(j['path'])
        exists = (out_dir / j['path']).exists()
        if args.force or not exists or not rec or rec.get('hash') != j['hash']:
            todo.append(j)
    print(f'  生成対象 {len(todo)} 件（生成済みで変更なし {len(jobs) - len(todo)} 件）')

    if args.dry_run:
        for j in todo:
            print(f'  {j["path"]:34} {j["voice"]:18} x{j["speed"]}  {j.get("text") or "（無音）"}')
        return
    if not todo:
        print('作り直しが必要な音声はありません。')
        return

    key = load_api_key()

    # 声の名前を先に確認しておく（途中で失敗すると中途半端になるため）
    needed = {(j['lang'], j['voice']) for j in todo if j['lang']}
    for lang in {l for l, _ in needed}:
        available = {v['name'] for v in list_voices(key, LANG_CODE[lang])}
        for l, name in needed:
            if l == lang and name not in available:
                sys.exit(f'声「{name}」が見つかりません。py generate_audio.py --list-voices '
                         f'{LANG_CODE[lang]} で使える名前を確認し、スクリプト冒頭の VOICES を直してください。')

    chars = sum(len(j.get('text') or '') for j in todo)
    print(f'  ビットレート：{MP3_BITRATE_KBPS}kbps')
    print(f'  送信する文字数の合計：約 {chars} 文字')

    failed = []
    for n, j in enumerate(todo, start=1):
        label = j.get('text') or '（無音）'
        print(f'[{n}/{len(todo)}] {j["path"]}  {label}')
        try:
            if j.get('silence_ms'):
                pcm = silence_pcm(j['silence_ms'])
            else:
                pcm = synthesize_pcm(key, j['voice'], j['lang'], j['speed'], j['text'])
            data, info = clean_mp3(encode_mp3(pcm, MP3_BITRATE_KBPS))
        except Exception as e:
            print(f'    失敗：{e}')
            failed.append(j['path'])
            continue
        dest = out_dir / j['path']
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(data)
        manifest['files'][j['path']] = {
            'id': j['id'], 'hash': j['hash'], 'text': j.get('text'), 'silenceMs': j.get('silence_ms'),
            'voice': j['voice'], 'speakingRate': j['speed'], **info,
        }
        manifest['generatedAt'] = datetime.now().isoformat(timespec='seconds')
        manifest_path.parent.mkdir(parents=True, exist_ok=True)
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=1), encoding='utf-8')
        if not j.get('silence_ms'):
            time.sleep(0.1)

    # ---- 連結できる状態かを点検 ----
    files = manifest['files']
    print('\n点検')
    srs = {f['sampleRate'] for f in files.values()}
    vers = {f['mpegVersion'] for f in files.values()}
    chs = {f['channelMode'] for f in files.values()}
    brs = {tuple(f['bitrateKbps']) for f in files.values()}
    ok = True
    if len(srs) > 1 or len(vers) > 1 or len(chs) > 1:
        ok = False
        print(f'  × 形式が揃っていません（周波数 {srs}、MPEG {vers}、チャンネル {chs}）。連結すると正しく再生できません。')
    not_cbr = [p for p, f in files.items() if not f['cbr']]
    if not_cbr:
        ok = False
        print(f'  × 可変ビットレートのファイルがあります（{len(not_cbr)} 件）。連結後の長さがずれる可能性があります。')
    if ok:
        print(f'  ○ すべて {srs.pop()} Hz・固定ビットレート {sorted(brs)} kbps で揃っています。そのまま連結できます。')

    known = {j['path'] for j in build_jobs(sentences, words)}
    orphans = sorted(p for p in files if p not in known)
    if orphans:
        print(f'  注意：Excel に無くなった文・単語の音声が残っています（自動では消しません）：')
        for p in orphans:
            print(f'    {p}')

    if failed:
        print(f'\n{len(failed)} 件が失敗しました。もう一度実行すると、失敗したものだけ作り直します。')
        sys.exit(1)
    print(f'\n完了：{out_dir.resolve()}')


if __name__ == '__main__':
    main()
